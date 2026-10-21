from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = "作品集内容采集表.docx"

INK = "141512"
BODY = "33352F"
MUTED = "777B72"
LINE = "D9DBD4"
SOFT = "F3F4EF"
GREEN = "5C711F"
SIGNAL = "C8F36C"
WHITE = "F3F1EA"
BLACK = "0A0B0B"

CN_FONT = "微软雅黑"
EN_FONT = "Arial"
CONTENT_WIDTH_CM = 16.6


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


def style_run(run, size=9.5, bold=False, color=BODY, italic=False, font=EN_FONT, eastasia=CN_FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), eastasia)
    rfonts.set(qn("w:ascii"), font)
    rfonts.set(qn("w:hAnsi"), font)
    return run


def add_run(paragraph, text, **kwargs):
    return style_run(paragraph.add_run(text), **kwargs)


P_BORDER_SUCCESSORS = (
    "w:shd",
    "w:tabs",
    "w:spacing",
    "w:ind",
    "w:contextualSpacing",
    "w:mirrorIndents",
    "w:suppressOverlap",
    "w:jc",
    "w:textDirection",
    "w:textAlignment",
    "w:textboxTightWrap",
    "w:outlineLvl",
    "w:divId",
    "w:cnfStyle",
    "w:rPr",
    "w:sectPr",
    "w:pPrChange",
)

TC_BORDER_SUCCESSORS = (
    "w:shd",
    "w:noWrap",
    "w:tcMar",
    "w:textDirection",
    "w:tcFitText",
    "w:vAlign",
    "w:hideMark",
)


def insert_before(container, element, successors):
    for child in container:
        if child.tag in {qn(tag) for tag in successors}:
            child.addprevious(element)
            return
    container.append(element)


def set_paragraph_border(paragraph, edge="bottom", color=LINE, size=4, space=4):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        insert_before(ppr, pbdr, P_BORDER_SUCCESSORS)
    existing = pbdr.find(qn(f"w:{edge}"))
    if existing is not None:
        pbdr.remove(existing)
    element = OxmlElement(f"w:{edge}")
    element.set(qn("w:val"), "single")
    element.set(qn("w:sz"), str(size))
    element.set(qn("w:space"), str(space))
    element.set(qn("w:color"), color)
    pbdr.append(element)


def shade_paragraph(paragraph, fill):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    insert_before(ppr, shd, P_BORDER_SUCCESSORS[1:])


def shade_cell(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    existing = tcpr.find(qn("w:shd"))
    if existing is not None:
        tcpr.remove(existing)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    insert_before(tcpr, shd, TC_BORDER_SUCCESSORS[1:])


def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcpr = cell._tc.get_or_add_tcPr()
    existing = tcpr.find(qn("w:tcMar"))
    if existing is not None:
        tcpr.remove(existing)
    tcmar = OxmlElement("w:tcMar")
    for name, value in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        el = OxmlElement(f"w:{name}")
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        tcmar.append(el)
    insert_before(tcpr, tcmar, TC_BORDER_SUCCESSORS[3:])


def clear_cell_borders(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    existing = tcpr.find(qn("w:tcBorders"))
    if existing is not None:
        tcpr.remove(existing)
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    insert_before(tcpr, borders, TC_BORDER_SUCCESSORS)


def remove_table_borders(table):
    tblpr = table._tbl.tblPr
    existing = tblpr.find(qn("w:tblBorders"))
    if existing is not None:
        tblpr.remove(existing)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    for child in tblpr:
        if child.tag in {qn("w:shd"), qn("w:tblLayout"), qn("w:tblCellMar"), qn("w:tblLook")}:
            child.addprevious(borders)
            break
    else:
        tblpr.append(borders)
    for row in table.rows:
        for cell in row.cells:
            clear_cell_borders(cell)


def set_cell_border(cell, edge="bottom", color=LINE, size=4):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        insert_before(tcpr, borders, TC_BORDER_SUCCESSORS)
    existing = borders.find(qn(f"w:{edge}"))
    if existing is not None:
        borders.remove(existing)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), "0")
    el.set(qn("w:color"), color)
    borders.append(el)


def set_col_width(cell, width_cm):
    cell.width = Cm(width_cm)
    tcpr = cell._tc.get_or_add_tcPr()
    tcw = tcpr.find(qn("w:tcW"))
    if tcw is None:
        tcw = OxmlElement("w:tcW")
        insert_before(tcpr, tcw, ("w:gridSpan", "w:hMerge", "w:vMerge", "w:tcBorders", "w:shd") + TC_BORDER_SUCCESSORS)
    tcw.set(qn("w:w"), str(int(width_cm * 567)))
    tcw.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.autofit = False
    for idx, width in enumerate(widths):
        try:
            table.columns[idx].width = Cm(width)
        except IndexError:
            pass
    for row in table.rows:
        for idx, width in enumerate(widths):
            if idx < len(row.cells):
                set_col_width(row.cells[idx], width)


def set_row_cant_split(row):
    trpr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    trpr.append(cant_split)


def add_page_field(paragraph, instruction):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t")
    fld_text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_text)
    run._r.append(fld_end)
    style_run(run, size=7.5, color=MUTED)


def add_header_footer(section):
    section.different_first_page_header_footer = True

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.paragraph_format.space_after = Pt(8)
    add_run(hp, "PORTFOLIO / CONTENT BRIEF", size=7, color=MUTED)
    add_run(hp, "  /  2026", size=7, color=MUTED)
    set_paragraph_border(hp, color=LINE, size=4, space=6)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(6)
    add_run(fp, "PAGE  ", size=7, color=MUTED)
    add_page_field(fp, "PAGE")
    add_run(fp, "  /  ", size=7, color=MUTED)
    add_page_field(fp, "NUMPAGES")


def add_cover(doc):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    remove_table_borders(table)
    row = table.rows[0]
    row.height = Cm(10.6)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    cell = row.cells[0]
    set_col_width(cell, CONTENT_WIDTH_CM)
    shade_cell(cell, BLACK)
    set_cell_margins(cell, top=700, bottom=700, left=760, right=760)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(60)
    add_run(p, "PORTFOLIO  /  CONTENT BRIEF", size=8, color="9FA49A")

    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    add_run(p, "个人作品集", size=34, bold=True, color=WHITE)

    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(22)
    add_run(p, "内容采集清单", size=27, bold=False, color=WHITE)

    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    set_paragraph_border(p, color=SIGNAL, size=18, space=2)
    add_run(p, "", size=1)

    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(0)
    add_run(p, "把经历、项目与审美，整理成一份可被看见的作品集。", size=9.5, color="A7ABA1")

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(22)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    add_run(p, "PORTFOLIO CONTENT BRIEF", size=8, bold=True, color=GREEN)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    add_run(p, "请尽可能完整地填写以下信息。", size=18, bold=True, color=INK)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.7
    p.paragraph_format.space_after = Pt(26)
    add_run(
        p,
        "这份清单用于梳理个人定位、经历、项目与视觉偏好。\n"
        "标有 ★ 的字段为上线必需项；不确定的内容可以填写“待定”。\n"
        "完成后请连同图片、视频与简历 PDF 一起提交。",
        size=9.5,
        color=BODY,
    )

    # Small note box
    note = doc.add_table(rows=1, cols=1)
    note.alignment = WD_TABLE_ALIGNMENT.CENTER
    note.autofit = False
    remove_table_borders(note)
    ncell = note.rows[0].cells[0]
    set_table_widths(note, [CONTENT_WIDTH_CM])
    shade_cell(ncell, SOFT)
    set_cell_margins(ncell, top=200, bottom=200, left=260, right=260)
    np = ncell.paragraphs[0]
    np.paragraph_format.line_spacing = 1.5
    add_run(np, "当前网站为演示版本。", size=9, bold=True, color=INK)
    add_run(np, " 页面中的姓名、经历与项目均为占位内容，本表用于收集正式信息并替换。", size=9, color=BODY)

    # Cover metadata
    info = doc.add_table(rows=1, cols=3)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    info.autofit = False
    remove_table_borders(info)
    widths = [5.5, 6.1, 5.0]
    values = [("版本", "V.01 / 2026"), ("适用项目", "个人作品集网站"), ("预计填写", "35–45 分钟")]
    set_table_widths(info, widths)
    for idx, (label, value) in enumerate(values):
        cell = info.rows[0].cells[idx]
        set_cell_margins(cell, top=180, bottom=40, left=0 if idx == 0 else 140, right=120)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        add_run(p, label, size=7, color=MUTED)
        p2 = cell.add_paragraph()
        add_run(p2, value, size=9, color=INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(26)
    add_run(p, "DESIGNED FOR YOUR PORTFOLIO / 2026", size=7, color=MUTED)


def add_contents(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    add_run(p, "CONTENTS", size=8, bold=True, color=GREEN)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    add_run(p, "目录", size=25, bold=True, color=INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(20)
    add_run(p, "填写顺序可以跳选；建议先完成基础信息与项目，再补充视觉偏好。", size=9, color=MUTED)

    rows = [
        ("01", "基础信息", "姓名、定位、城市、个人视觉身份"),
        ("02", "联系与社交", "邮箱、微信、平台链接与公开范围"),
        ("03", "定位与文案", "首屏标题、个人介绍、语气与关键词"),
        ("04", "数据与经历", "项目数据、工作经历、教育与奖项"),
        ("05", "精选项目", "4 个项目的背景、策略、成果与素材"),
        ("06", "能力与优势", "5 项核心能力、说明与标签"),
        ("07", "视觉偏好", "参考网站、色彩、字体、动效与影像"),
        ("08", "上线与素材", "域名、语言、版权、文件清单"),
        ("09", "补充与确认", "待定项、保密范围与确认签名"),
    ]
    table = doc.add_table(rows=len(rows), cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    remove_table_borders(table)
    widths = [1.2, 4.9, 10.5]
    set_table_widths(table, widths)
    for i, (num, title, desc) in enumerate(rows):
        cells = table.rows[i].cells
        for j in range(3):
            set_cell_margins(cells[j], top=100, bottom=100, left=0 if j == 0 else 120, right=80)
            set_cell_border(cells[j], "bottom", LINE, 4)
        p = cells[0].paragraphs[0]
        add_run(p, num, size=8, bold=True, color=GREEN)
        p = cells[1].paragraphs[0]
        add_run(p, title, size=10.5, bold=True, color=INK)
        p = cells[2].paragraphs[0]
        add_run(p, desc, size=8.5, color=MUTED)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(8)
    add_run(p, "填写说明 / HOW TO USE", size=8, bold=True, color=GREEN)

    notes = [
        ("直接在字段后填写。", "所有字段均可编辑；长文本会自动换行。"),
        ("先写事实，再写形容词。", "项目背景、动作、结果尽量具体；数字比形容词更有说服力。"),
        ("素材另附。", "图片、视频、Logo 不需要嵌入 Word；按“项目编号-类型-名称”命名后打包。"),
        ("不确定或保密，请标注。", "填写“待定”“需隐去客户名”“不可公开”即可，不要直接留空。"),
    ]
    for title, desc in notes:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.25)
        p.paragraph_format.space_after = Pt(6)
        add_run(p, "—  ", size=9, color=GREEN)
        add_run(p, title, size=9, bold=True, color=INK)
        add_run(p, "  " + desc, size=9, color=BODY)


def add_section_heading(doc, num, title, en, intro, page_break=False):
    if page_break:
        doc.add_page_break()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20 if not page_break else 0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    add_run(p, f"SECTION {num}  /  {en}", size=7.5, bold=True, color=GREEN)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.tab_stops.add_tab_stop(Cm(CONTENT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT)
    add_run(p, num, size=10, bold=True, color=GREEN)
    add_run(p, "   " + title, size=20, bold=True, color=INK)
    add_run(p, "\t" + en, size=7.5, bold=True, color=MUTED)
    set_paragraph_border(p, "bottom", INK, 10, 6)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.6
    p.paragraph_format.keep_with_next = True
    add_run(p, intro, size=8.8, color=MUTED)


def add_field(doc, label, hint=None, options=None, required=False, tall=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7 if tall else 4)
    p.paragraph_format.space_after = Pt(5 if tall else 3)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.keep_together = True
    if required:
        add_run(p, "★ ", size=8, bold=True, color=GREEN)
    add_run(p, label + "：", size=10, bold=True, color=INK)
    if hint:
        add_run(p, f"  {hint}", size=7.5, color=MUTED)
    if options:
        add_run(p, "  " + options, size=9, color=BODY)
    set_paragraph_border(p, "bottom", LINE, 4, 5)
    return p


def add_group(doc, text, note=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(13)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    shade_paragraph(p, SOFT)
    add_run(p, "  " + text, size=8.5, bold=True, color=INK)
    if note:
        add_run(p, "    " + note, size=7, color=MUTED)
    set_paragraph_border(p, "left", GREEN, 12, 4)
    return p


def add_checklist_table(doc):
    pairs = [
        ("个人照片（4:5，≥2000px）", "Hero 视频 / 首屏主视觉"),
        ("项目封面图（每项目 1 张）", "项目细节图（每项目 3–6 张）"),
        ("个人 Logo（SVG / PNG）", "客户 / 品牌 Logo"),
        ("个人简历 PDF", "字体与图片授权"),
        ("参考截图 / 灵感板", "其他素材"),
    ]
    table = doc.add_table(rows=len(pairs) + 1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    remove_table_borders(table)
    widths = [5.1, 3.2, 5.1, 3.2]
    set_table_widths(table, widths)

    headers = ["素材 / 文件", "状态", "素材 / 文件", "状态"]
    header_row = table.rows[0]
    set_row_cant_split(header_row)
    for j, header in enumerate(headers):
        cell = header_row.cells[j]
        shade_cell(cell, INK)
        set_cell_margins(cell, top=80, bottom=80, left=90, right=90)
        p = cell.paragraphs[0]
        add_run(p, header, size=7.5, bold=True, color=WHITE)

    for i, pair in enumerate(pairs, start=1):
        row = table.rows[i]
        set_row_cant_split(row)
        cells = row.cells
        for j in range(4):
            set_cell_margins(cells[j], top=65, bottom=65, left=90, right=90)
            set_cell_border(cells[j], "bottom", LINE, 4)
            # keep the checklist on one page where possible
            for paragraph in cells[j].paragraphs:
                paragraph.paragraph_format.keep_with_next = i < len(pairs)
        p = cells[0].paragraphs[0]
        add_run(p, pair[0], size=8.2, color=BODY)
        p = cells[1].paragraphs[0]
        add_run(p, "□ 已有    □ 待补", size=7, color=MUTED)
        p = cells[2].paragraphs[0]
        add_run(p, pair[1], size=8.2, color=BODY)
        p = cells[3].paragraphs[0]
        add_run(p, "□ 已有    □ 待补", size=7, color=MUTED)


def add_final_note(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.6
    shade_paragraph(p, SOFT)
    add_run(p, "  " + text, size=8.5, color=MUTED)
    set_paragraph_border(p, "left", GREEN, 12, 4)


def build_document():
    doc = Document()

    # Page and base styles
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(0.9)
    section.footer_distance = Cm(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = EN_FONT
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = rgb(BODY)
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_after = Pt(0)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), CN_FONT)
    rfonts.set(qn("w:ascii"), EN_FONT)
    rfonts.set(qn("w:hAnsi"), EN_FONT)

    add_header_footer(section)

    props = doc.core_properties
    props.title = "个人作品集内容采集表"
    props.subject = "Portfolio Content Brief"
    props.author = "Portfolio Studio"
    props.keywords = "作品集, 内容采集, 品牌设计, 视觉设计, AI 设计"

    # Cover + contents
    add_cover(doc)
    doc.add_page_break()
    add_contents(doc)
    doc.add_page_break()

    # 01 Identity
    add_section_heading(
        doc,
        "01",
        "基础信息",
        "IDENTITY",
        "先确认你在网站上的“身份”。这里的信息会出现在导航、首屏、关于我与页脚中。",
    )
    add_field(doc, "中文姓名", "必填", required=True)
    add_field(doc, "英文名 / 拼音", "用于 Logo、页脚与英文标题")
    add_field(doc, "品牌 / 工作室名称", "如以个人名义展示可留空")
    add_field(doc, "个人标识字母", "例如 LX、XY；用于导航与页脚 Logo")
    add_field(doc, "职位定位（中文）", "建议 2–3 个", required=True)
    add_field(doc, "职位定位（英文）", "例如 Visual / AI / Brand Designer")
    add_field(doc, "所在城市 / 国家", "必填", required=True)
    add_field(doc, "希望被如何称呼", "姓名 / 英文名 / 老师 / 其他")
    add_field(
        doc,
        "目前状态",
        options="□ 可接新项目    □ 档期有限    □ 暂不接项目",
    )
    add_field(doc, "可合作档期", "例如 2026 Q1 起 / 每月 1–2 个项目")
    add_field(doc, "工作语言", "例如 中文 / 英文 / 中英双语")
    add_field(
        doc,
        "个人照片状态",
        options="□ 已附    □ 已有但需挑选    □ 需要拍摄    □ 用插画 / 抽象图代替",
    )
    add_field(
        doc,
        "是否公开真实姓名",
        options="□ 公开中文名    □ 仅公开英文名    □ 使用化名",
    )
    add_field(doc, "个人 Logo 文件", options="□ 已附 SVG    □ 已附 PNG    □ 仅使用文字标    □ 待设计")

    # 02 Contact
    add_section_heading(
        doc,
        "02",
        "联系与社交",
        "CONTACT",
        "用于导航 CTA、关于我与整屏联系页。请注明哪些信息可以公开。",
    )
    add_field(doc, "公开邮箱", "必填", required=True)
    add_field(doc, "商务电话", "如不希望公开可留空")
    add_field(doc, "微信", "微信号 / 二维码")
    add_field(doc, "小红书", "主页链接或账号名")
    add_field(doc, "Behance", "主页链接")
    add_field(doc, "站酷", "主页链接")
    add_field(doc, "Instagram", "主页链接")
    add_field(doc, "LinkedIn", "主页链接")
    add_field(doc, "个人网站 / 其他平台", "可填写多个")
    add_field(
        doc,
        "希望公开展示的联系方式",
        options="□ 邮箱    □ 微信    □ 电话    □ 社交平台    □ 仅邮箱",
    )
    add_field(doc, "是否公开手机号", options="□ 是    □ 否")
    add_field(doc, "回复时间说明", "例如 通常在 24 小时内回复")

    # 03 Positioning / copy
    add_section_heading(
        doc,
        "03",
        "定位与文案",
        "POSITIONING & COPY",
        "这一部分决定网站的语气。短期写不完整也没关系，先把真实想法写下来，我可以据此提炼。",
    )
    add_field(doc, "一句话定位", "20 字以内，必填", required=True)
    add_field(doc, "首屏主标题建议", "6–12 字；可留空由我代写")
    add_field(doc, "首屏英文副标题", "一句短句，例如 A practice of visual order")
    add_field(doc, "首屏简介", "1–2 句，建议 40–80 字", tall=True)
    add_field(doc, "你希望访客第一眼记住什么", "关键词或一句话", tall=True)
    add_field(doc, "你与同类设计师最大的不同", "能力组合 / 工作方式 / 审美判断", tall=True)
    add_field(doc, "设计理念", "50–100 字", tall=True)
    add_field(doc, "100 字个人介绍", "用于模块引言", tall=True)
    add_field(doc, "200 字个人介绍", "用于 About 模块完整介绍", tall=True)
    add_field(doc, "三个核心标签", "例如 品牌系统 / AI 创意 / 视觉叙事")
    add_field(doc, "服务范围", "例如 品牌识别、艺术指导、数字体验", tall=True)
    add_field(doc, "目标客户 / 理想合作方", "行业、规模、合作方式", tall=True)
    add_field(doc, "你最想接的项目", "类型 / 行业 / 合作深度", tall=True)
    add_field(doc, "你不希望接的项目", "避免后期沟通错位", tall=True)
    add_field(doc, "联系页邀请文案", "1–2 句，用于 LET'S TALK 区域", tall=True)
    add_field(
        doc,
        "文案语气",
        options="□ 克制专业    □ 文艺感性    □ 直接有力    □ 技术理性    □ 其他",
    )
    add_field(doc, "必须出现的关键词", "3–8 个")
    add_field(doc, "不希望出现的词 / 表达", "例如“赋能”“闭环”“顶级”")
    add_field(doc, "未来 1–2 年的方向", "有助于判断作品与经历的取舍", tall=True)

    # 04 Data & experience
    add_section_heading(
        doc,
        "04",
        "数据与经历",
        "DATA & EXPERIENCE",
        "先写事实，再决定展示顺序。所有数字都可以标注“约”或“保密”。",
    )
    add_field(doc, "工作年限", "例如 8 年")
    add_field(doc, "完成项目数量", "例如 46 个")
    add_field(doc, "合作行业数量", "例如 12 个")
    add_field(doc, "获得奖项数量", "例如 9 项")
    add_field(doc, "服务客户数量", "可选")
    add_field(doc, "其他可量化数据", "例如 内容曝光、转化提升、团队规模")
    add_field(doc, "可公开的客户 / 品牌", "用顿号分隔", tall=True)

    for idx in range(1, 4):
        add_group(doc, f"经历 {idx:02d}", "按时间倒序排列" if idx == 1 else None)
        add_field(doc, "时间", "例如 2023 — NOW")
        add_field(doc, "公司 / 机构")
        add_field(doc, "职位")
        add_field(doc, "地点", "城市 / 远程")
        add_field(doc, "负责内容", "用动词开头，1–3 句", tall=True)
        add_field(doc, "关键成果", "数据、规模、代表性项目", tall=True)
        add_field(doc, "是否在网站展示", options="□ 是    □ 否")

    add_group(doc, "教育背景")
    add_field(doc, "时间")
    add_field(doc, "学校")
    add_field(doc, "专业 / 学位")
    add_field(doc, "补充说明", "可选")

    for idx in range(1, 4):
        add_group(doc, f"奖项 / 展览 / 媒体 {idx:02d}")
        add_field(doc, "年份")
        add_field(doc, "名称")
        add_field(doc, "颁发机构 / 媒体")
        add_field(doc, "链接或说明")

    add_final_note(doc, "如果经历较多，请优先提供与“品牌 / AI / 视觉 / 数字体验”最相关的 3–5 段。")

    add_group(doc, "经历写作提示 / WRITING NOTES")
    for line in [
        "每段经历建议写清三件事：职责范围、代表性项目、可量化结果。",
        "用动词开头：主导、建立、重构、推动、交付，比“负责”更有信息量。",
        "客户或项目不能公开时，可写成“某头部消费电子品牌”“某国际香氛品牌”。",
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.2)
        p.paragraph_format.space_after = Pt(5)
        add_run(p, "—  ", size=8.5, color=GREEN)
        add_run(p, line, size=8.5, color=BODY)

    # 05 Projects
    add_section_heading(
        doc,
        "05",
        "精选项目",
        "SELECTED PROJECTS",
        "建议提供 4 个项目，至少 3 个。每个项目重点说明：为什么做、你怎么想、最后改变了什么。\n"
        "长文本建议控制在 50–150 字；图片不需要嵌入 Word，填写文件名即可。",
        page_break=True,
    )
    for idx in range(1, 5):
        add_group(doc, f"项目 {idx:02d}  /  PROJECT {idx:02d}", "请完整填写后再进入下一个")
        add_field(doc, "项目名称（中文）", required=True)
        add_field(doc, "项目名称（英文）")
        add_field(
            doc,
            "项目类型",
            options="□ 品牌    □ AI 创意    □ 数字体验    □ Campaign    □ 包装    □ 其他",
        )
        add_field(doc, "项目时间", "年份或月份")
        add_field(doc, "客户 / 行业", "如不便公开请注明")
        add_field(doc, "我的角色", "例如 创意总监 / 品牌设计 / 独立完成")
        add_field(doc, "项目背景 / 挑战", "50–100 字", tall=True)
        add_field(doc, "设计策略 / 概念", "50–100 字", tall=True)
        add_field(doc, "最终方案", "100–150 字，建议写清视觉系统与关键决策", tall=True)
        add_field(doc, "项目成果", "数据、客户反馈、上线情况", tall=True)
        add_field(doc, "交付内容", "例如 品牌识别、包装、动态、网站")
        add_field(
            doc,
            "可公开程度",
            options="□ 完整公开    □ 仅展示图片    □ 隐去客户名    □ 保密",
        )
        add_field(doc, "项目链接", "Behance / 网站 / 视频链接")
        add_field(doc, "封面图文件名", "例如 01-cover-aether.jpg")
        add_field(doc, "其他图片文件名", "3–6 张，用顿号分隔")
        add_field(doc, "是否允许首页展示", options="□ 是    □ 否")

    add_final_note(doc, "如果项目超过 4 个，请复制“项目 04”整组字段继续填写，并标注优先级。")

    # 06 Capabilities
    add_section_heading(
        doc,
        "06",
        "能力与优势",
        "CAPABILITIES",
        "网站会用 5 张卡片呈现你的能力。请避免“全能型”描述，尽量说明你如何解决问题。",
    )
    for idx in range(1, 6):
        add_group(doc, f"能力 {idx:02d}")
        add_field(doc, "能力名称（中文）", required=True)
        add_field(doc, "英文名称")
        add_field(doc, "一句话说明", "12–20 字")
        add_field(doc, "详细说明", "40–80 字，写清能力带来的结果", tall=True)
        add_field(doc, "关键词标签", "3–4 个，用顿号分隔")
        add_field(doc, "对应项目 / 案例", "可填写项目编号")

    # 07 Visual preferences
    add_section_heading(
        doc,
        "07",
        "视觉偏好",
        "VISUAL DIRECTION",
        "这一部分帮助我判断什么“像你”。可以只写关键词，也可以直接给参考链接与截图。",
    )
    add_field(doc, "网站整体关键词", "3–5 个，例如 克制、冷峻、未来感", tall=True)
    add_field(doc, "喜欢的参考网站 01", "链接 + 你喜欢的具体一点", tall=True)
    add_field(doc, "喜欢的参考网站 02", "链接 + 你喜欢的具体一点", tall=True)
    add_field(doc, "喜欢的参考网站 03", "链接 + 你喜欢的具体一点", tall=True)
    add_field(doc, "不喜欢的风格 / 避免", "例如 模板感、霓虹赛博、过度动效", tall=True)
    add_field(doc, "希望保留的视觉元素", "例如 液态金属、网格、胶片颗粒", tall=True)
    add_field(doc, "主色偏好", "可写颜色名或色值")
    add_field(doc, "辅助色 / 强调色", "可写颜色名或色值")
    add_field(doc, "字体偏好", "装饰体 / 无衬线 / 衬线 / 中英混排")
    add_field(doc, "动效程度", options="□ 克制    □ 适中    □ 强烈")
    add_field(
        doc,
        "影像气质",
        options="□ 黑白    □ 冷色    □ 暖色    □ 高对比    □ 彩色    □ 混合",
    )
    add_field(doc, "是否需要视频背景", options="□ 是    □ 否    □ 可讨论")
    add_field(doc, "是否需要中英双语", options="□ 是    □ 否    □ 未来扩展")
    add_field(doc, "项目图片风格", options="□ 摄影    □ 3D / 抽象    □ 排版海报    □ 混合")
    add_field(doc, "参考截图", "请另附文件，并注明对应模块")
    add_field(doc, "其他视觉要求", tall=True)

    add_group(doc, "写作提示 / WRITING NOTES", "不擅长写文案时，可以直接按这些结构回答")
    for line in [
        "项目描述四步：背景（为什么做）→ 挑战（难在哪）→ 行动（你做了什么）→ 结果（改变了什么）。",
        "用具体名词和数字代替形容词。与其写“提升了品牌调性”，不如写“统一了 12 个触点的视觉规则”。",
        "一句话定位的结构：为【谁】提供【什么】，通过【方法】实现【结果】。",
        "每个项目只保留一个最值得讲的决策，不要写成项目流水账。",
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.2)
        p.paragraph_format.space_after = Pt(5)
        add_run(p, "—  ", size=8.5, color=GREEN)
        add_run(p, line, size=8.5, color=BODY)

    # 08 Launch & assets
    add_section_heading(
        doc,
        "08",
        "上线与素材",
        "LAUNCH & ASSETS",
        "用于确认网站边界、上线方式与素材版权。没有确定答案的项目可以写“待讨论”。",
        page_break=True,
    )
    add_field(doc, "网站形式", options="□ 单页作品集    □ 多页    □ 可讨论")
    add_field(doc, "域名", "已有 / 待注册 / 使用临时地址")
    add_field(doc, "部署平台偏好", "Vercel / Netlify / 自有服务器 / 不确定")
    add_field(doc, "是否需要后台 / CMS", options="□ 是    □ 否    □ 未来考虑")
    add_field(doc, "是否需要简历下载", options="□ 是    □ 否")
    add_field(doc, "是否需要联系表单", options="□ 是    □ 否（仅邮箱）")
    add_field(doc, "是否展示微信二维码", options="□ 是    □ 否")
    add_field(doc, "是否需要英文版", options="□ 是    □ 否    □ 未来扩展")
    add_field(doc, "是否需要项目详情页", options="□ 是    □ 否    □ 弹窗即可")
    add_field(doc, "版权署名格式", "例如 © 2026 林序 — All Rights Reserved")
    add_field(doc, "备案 / 合规要求", "国内备案、隐私政策、Cookie 等")
    add_field(doc, "素材交付方式", "网盘链接 / 邮件 / 微信文件")

    add_group(doc, "素材清单", "不需要嵌入 Word；按文件名整理后打包即可")
    add_checklist_table(doc)

    add_group(doc, "素材命名建议")
    add_field(doc, "统一命名格式", "项目编号-类型-名称，例如 01-cover-aether.jpg", tall=True)
    add_field(doc, "图片尺寸 / 格式", "JPG / PNG / WebP；宽度 ≥1600px")
    add_field(doc, "视频尺寸 / 时长", "16:9；8–15 秒；MP4 / WebM")
    add_field(doc, "其他补充", tall=True)

    # 09 Final confirmation
    add_section_heading(
        doc,
        "09",
        "补充与确认",
        "FINAL CHECK",
        "填写完以上内容后，请在这里标注仍需补充的事项，并确认可公开范围。",
    )
    add_field(doc, "仍需补充 / 待定事项", "我会根据这里继续向你追问", tall=True)
    add_field(doc, "不能公开的信息", "姓名、客户、数据、项目图等", tall=True)
    add_field(doc, "需要隐去的客户 / 项目", "请具体说明", tall=True)
    add_field(doc, "可公开的客户 Logo", "用顿号分隔", tall=True)
    add_field(doc, "其他要求 / 备注", tall=True)
    add_field(doc, "期望上线时间", "例如 2026 年 3 月")
    add_field(doc, "方便沟通的时间", "工作日 / 晚间 / 周末")
    add_field(
        doc,
        "公开确认",
        options="□ 以上内容可用于网站制作    □ 部分内容需二次确认",
    )
    add_field(doc, "填写人 / 日期", "签名或姓名 + 日期", tall=True)

    add_final_note(doc, "完成后请将本文件与素材包一起发回。收到后我会先整理信息结构，再进入视觉替换与页面优化。")

    # Force field updates on open
    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)

    doc.save(OUT)
    print(f"saved {OUT}")


if __name__ == "__main__":
    build_document()
